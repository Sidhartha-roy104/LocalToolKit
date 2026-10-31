"""
Automated Smoke Test Suite for LocalToolkit.
Tests all endpoints, tools, security guards, settings, and static offline requirements.
"""
from __future__ import annotations
import base64
import io
import json
import os
import re
import sys
import urllib.parse
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw
import fitz  # PyMuPDF
import docx

from backend.main import app, LAUNCH_TOKEN
from backend.utils.deps import get_deps
from backend.tools.user_settings import get_default_my_files_dir, load_settings, save_settings


# ── Sample File Generation Helpers ──────────────────────────────────────────

def create_sample_jpg_photo() -> bytes:
    """Return a 400x500 RGB photo with a simple drawn face and EXIF metadata."""
    img = Image.new("RGB", (400, 500), color=(240, 240, 245))
    draw = ImageDraw.Draw(img)
    # Draw head
    draw.ellipse((100, 80, 300, 280), fill=(255, 218, 185), outline=(180, 140, 100), width=3)
    # Draw eyes
    draw.ellipse((145, 150, 175, 175), fill=(40, 40, 40))
    draw.ellipse((225, 150, 255, 175), fill=(40, 40, 40))
    # Draw smile
    draw.arc((160, 190, 240, 230), start=0, end=180, fill=(180, 50, 50), width=4)
    # Draw shoulders
    draw.polygon([(40, 500), (360, 500), (320, 340), (80, 340)], fill=(50, 90, 160))

    # Add EXIF metadata to test removal
    exif = img.getexif()
    exif[0x010E] = "LocalToolkit Smoke Test Photo"
    exif[0x010F] = "TestCamera"
    exif[0x0110] = "Model X"

    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=90, exif=exif)
    return buf.getvalue()


def create_sample_png_alpha() -> bytes:
    """Return a 300x300 RGBA PNG with transparent sections."""
    img = Image.new("RGBA", (300, 300), color=(0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    draw.ellipse((30, 30, 270, 270), fill=(34, 197, 94, 220), outline=(21, 128, 61, 255), width=4)
    draw.rectangle((80, 130, 220, 170), fill=(255, 255, 255, 240))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def create_sample_cmyk_jpeg() -> bytes:
    """Return a 300x300 CMYK JPEG."""
    img = Image.new("CMYK", (300, 300), color=(100, 50, 0, 20))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def create_sample_signature_png() -> bytes:
    """Return a 300x120 RGBA transparent signature."""
    img = Image.new("RGBA", (300, 120), color=(255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    # Draw signature curves
    draw.line([(30, 70), (80, 40), (120, 80), (160, 30), (220, 70), (270, 50)], fill=(0, 20, 80, 255), width=3)
    draw.line([(60, 90), (240, 90)], fill=(0, 20, 80, 255), width=2)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def create_sample_3page_pdf() -> bytes:
    """Return a 3-page PDF with text, an RGB image, and an image with soft mask (SMask)."""
    doc = fitz.open()

    # Page 1: Standard text + RGB image
    p1 = doc.new_page(width=595, height=842)
    p1.insert_text((50, 72), "Document Title - Page 1", fontsize=18)
    p1.insert_text((50, 100), "This is confidential sample text for LocalToolkit testing.", fontsize=11)
    rgb_img = Image.new("RGB", (250, 180), (70, 130, 180))
    b1 = io.BytesIO()
    rgb_img.save(b1, format="JPEG", quality=85)
    p1.insert_image(fitz.Rect(50, 130, 300, 310), stream=b1.getvalue())

    # Page 2: Soft-masked image (alpha transparency)
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((50, 72), "Soft-Masked Image Test - Page 2", fontsize=18)
    p2.insert_text((50, 100), "Sensitive text: SECRET_AADHAAR_9999 to test redaction.", fontsize=12)
    rgba_img = Image.new("RGBA", (200, 200), (239, 68, 68, 180))
    b2 = io.BytesIO()
    rgba_img.save(b2, format="PNG")
    p2.insert_image(fitz.Rect(50, 140, 250, 340), stream=b2.getvalue())

    # Page 3: Additional text
    p3 = doc.new_page(width=595, height=842)
    p3.insert_text((50, 72), "Conclusion & Signatures - Page 3", fontsize=18)
    p3.insert_text((50, 120), "Signatures and verification stamps appear below.", fontsize=12)

    data = doc.tobytes()
    doc.close()
    return data


def create_sample_120page_pdf() -> bytes:
    """Return a 120-page PDF generated fast."""
    doc = fitz.open()
    for i in range(120):
        p = doc.new_page(width=595, height=842)
        p.insert_text((50, 80), f"LocalToolkit Large Document - Page {i + 1} of 120", fontsize=14)
        p.insert_text((50, 120), f"Content line for page index {i + 1}.", fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def create_sample_docx() -> bytes:
    """Return a minimal valid DOCX file."""
    doc = docx.Document()
    doc.add_heading("LocalToolkit Test Document", level=0)
    doc.add_paragraph("This is a paragraph created for smoke testing Word conversion.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture(scope="session")
def client():
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="session")
def jpg_photo():
    return create_sample_jpg_photo()


@pytest.fixture(scope="session")
def png_alpha():
    return create_sample_png_alpha()


@pytest.fixture(scope="session")
def cmyk_jpeg():
    return create_sample_cmyk_jpeg()


@pytest.fixture(scope="session")
def sig_png():
    return create_sample_signature_png()


@pytest.fixture(scope="session")
def pdf_3page():
    return create_sample_3page_pdf()


@pytest.fixture(scope="session")
def pdf_120page():
    return create_sample_120page_pdf()


@pytest.fixture(scope="session")
def sample_docx():
    return create_sample_docx()


# ── STEP 1: PDF Tools ────────────────────────────────────────────────────────

def test_pdf_merge(client, pdf_3page):
    files = [
        ("files", ("part1.pdf", pdf_3page, "application/pdf")),
        ("files", ("part2.pdf", pdf_3page, "application/pdf")),
    ]
    res = client.post("/api/pdf/merge", files=files)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 6
    doc.close()


def test_pdf_split(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"split_mode": "all", "page_ranges": ""}
    res = client.post("/api/pdf/split", files=files, data=data)
    assert res.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(res.content))
    namelist = zf.namelist()
    assert len(namelist) == 3


def test_pdf_extract(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"pages": "1,3"}
    res = client.post("/api/pdf/extract", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 2
    doc.close()


def test_pdf_remove(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"pages": "2"}
    res = client.post("/api/pdf/remove", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 2
    doc.close()


def test_pdf_organize(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"order": "3,2,1"}
    res = client.post("/api/pdf/organize", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 3
    # First page should now be the old page 3
    txt = doc[0].get_text()
    assert "Page 3" in txt
    doc.close()


def test_pdf_rotate(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"angle": "90", "pages": "all"}
    res = client.post("/api/pdf/rotate", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert doc[0].rotation == 90
    doc.close()


def test_pdf_pagenumbers(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"position": "bottom-center", "format": "page", "start_num": "1", "margin": "36"}
    res = client.post("/api/pdf/pagenumbers", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert "1" in doc[0].get_text()
    doc.close()


def test_pdf_watermark(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {
        "text": "CONFIDENTIAL",
        "opacity": "0.3",
        "rotation": "45",
        "fontsize": "48",
        "color": "gray",
    }
    res = client.post("/api/pdf/watermark", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 3
    doc.close()


def test_pdf_protect_and_unlock(client, pdf_3page):
    # 1. Protect
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"password": "testpassword123"}
    res_prot = client.post("/api/pdf/protect", files=files, data=data)
    assert res_prot.status_code == 200
    prot_doc = fitz.open(stream=res_prot.content, filetype="pdf")
    assert prot_doc.is_encrypted
    prot_doc.close()

    # 2. Unlock
    files_un = [("file", ("protected.pdf", res_prot.content, "application/pdf"))]
    data_un = {"password": "testpassword123"}
    res_un = client.post("/api/pdf/unlock", files=files_un, data=data_un)
    assert res_un.status_code == 200
    un_doc = fitz.open(stream=res_un.content, filetype="pdf")
    assert not un_doc.is_encrypted
    assert len(un_doc) == 3
    un_doc.close()


def test_pdf_to_images(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"fmt": "jpg", "dpi": "150"}
    res = client.post("/api/pdf/to-images", files=files, data=data)
    assert res.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(res.content))
    imgs = zf.namelist()
    assert len(imgs) == 3
    first_img_data = zf.read(imgs[0])
    img = Image.open(io.BytesIO(first_img_data))
    assert img.size[0] > 0 and img.size[1] > 0


def test_pdf_compress_and_no_black_blocks(client, pdf_3page):
    """Assert PDF compress executes, and output pages are NOT mostly black."""
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"quality": "medium", "target_kb": "0"}
    res = client.post("/api/pdf/compress", files=files, data=data)
    assert res.status_code == 200

    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 3

    # Check page rendering: page 2 contains soft-masked image, assert not corrupted to black blocks
    for page_idx in range(len(doc)):
        pix = doc[page_idx].get_pixmap()
        samples = pix.samples
        total_pixels = pix.width * pix.height
        # Count black pixels
        black_count = 0
        step = pix.n
        for i in range(0, len(samples), step):
            if samples[i] == 0 and samples[i+1] == 0 and samples[i+2] == 0:
                black_count += 1
        black_ratio = black_count / total_pixels
        assert black_ratio < 0.25, f"Page {page_idx+1} has {black_ratio:.1%} black pixels (black block corruption)"

    doc.close()


# ── STEP 2: Image Tools ─────────────────────────────────────────────────────

def test_image_compress_quality(client, jpg_photo):
    files = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data = {"quality": "70", "target_kb": "0"}
    res = client.post("/api/image/compress", files=files, data=data)
    assert res.status_code == 200
    img = Image.open(io.BytesIO(res.content))
    assert img.size == (400, 500)


def test_image_compress_target_kb_reachable(client, jpg_photo):
    files = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data = {"target_kb": "30", "quality": "75"}
    res = client.post("/api/image/compress", files=files, data=data)
    assert res.status_code == 200
    assert len(res.content) <= 30 * 1024 or res.headers.get("x-target-reachable") == "false"
    img = Image.open(io.BytesIO(res.content))
    assert img.size[0] > 0


def test_image_compress_target_kb_unreachable(client, jpg_photo):
    files = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data = {"target_kb": "1", "quality": "75"}
    res = client.post("/api/image/compress", files=files, data=data)
    assert res.status_code == 200
    # Meets target (<= 1024 bytes) or clearly indicates unreachable
    assert len(res.content) <= 1024 or res.headers.get("x-target-reachable") == "false" or "x-warning" in res.headers


def test_image_convert(client, jpg_photo):
    files = [("files", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data = {"to_format": "png"}
    res = client.post("/api/image/convert", files=files, data=data)
    assert res.status_code == 200
    img = Image.open(io.BytesIO(res.content))
    assert img.format == "PNG"


def test_image_to_pdf(client, jpg_photo):
    files = [("files", ("photo.jpg", jpg_photo, "image/jpeg"))]
    res = client.post("/api/image/to-pdf", files=files)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 1
    doc.close()


def test_image_resize_crop(client, jpg_photo):
    files = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data = {
        "width": "200",
        "height": "250",
        "fmt": "jpg",
        "quality": "85",
        "unit": "px",
        "dpi": "150",
    }
    res = client.post("/api/image/resize-crop", files=files, data=data)
    assert res.status_code == 200
    img = Image.open(io.BytesIO(res.content))
    assert img.size == (200, 250)


def test_estimate_size_endpoints(client, jpg_photo, pdf_3page):
    # 1. Image estimate-size
    files_img = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    data_img = {"quality": "60", "target_kb": "40"}
    res_img = client.post("/api/image/estimate-size", files=files_img, data=data_img)
    assert res_img.status_code == 200
    json_img = res_img.json()
    assert "estimated_bytes" in json_img
    assert json_img["estimated_bytes"] > 0

    # 2. PDF estimate-size
    files_pdf = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data_pdf = {"quality": "medium", "target_kb": "0", "dpi": "150"}
    res_pdf = client.post("/api/pdf/estimate-size", files=files_pdf, data=data_pdf)
    assert res_pdf.status_code == 200
    json_pdf = res_pdf.json()
    assert "estimated_bytes" in json_pdf
    assert json_pdf["estimated_bytes"] > 0


# ── STEP 3: Document Tools ──────────────────────────────────────────────────

def test_word_to_pdf_skip_or_run(client, sample_docx):
    deps = get_deps()
    if not deps.get("libreoffice"):
        pytest.skip("LibreOffice is not installed on this system (Word-to-PDF skipped as expected).")
    files = [("file", ("test.docx", sample_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))]
    res = client.post("/api/doc/word-to-pdf", files=files)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) >= 1
    doc.close()


def test_pdf_to_word(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/doc/pdf-to-word", files=files)
    assert res.status_code == 200
    doc = docx.Document(io.BytesIO(res.content))
    assert len(doc.paragraphs) > 0


# ── STEP 4: Extra & Govt Form Tools ─────────────────────────────────────────

def test_remove_metadata_image(client, jpg_photo):
    # Verify input photo had EXIF
    orig_img = Image.open(io.BytesIO(jpg_photo))
    assert len(orig_img.getexif()) > 0

    files = [("file", ("photo.jpg", jpg_photo, "image/jpeg"))]
    res = client.post("/api/tools/remove-metadata", files=files)
    assert res.status_code == 200

    clean_img = Image.open(io.BytesIO(res.content))
    assert len(clean_img.getexif()) == 0


def test_remove_metadata_pdf(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/tools/remove-metadata", files=files)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    meta = doc.metadata
    assert not meta.get("author")
    assert not meta.get("title")
    doc.close()


def test_combine_photo_signature(client, jpg_photo, sig_png):
    files = [
        ("photo", ("photo.jpg", jpg_photo, "image/jpeg")),
        ("signature", ("sig.png", sig_png, "image/png")),
    ]
    data = {
        "name_text": "RAHUL SHARMA",
        "date_text": "01/11/2026",
        "target_width": "350",
        "target_height": "500",
    }
    res = client.post("/api/tools/combine-photo-signature", files=files, data=data)
    assert res.status_code == 200
    img = Image.open(io.BytesIO(res.content))
    assert img.size == (350, 500)


def test_pdf_sign(client, pdf_3page, sig_png):
    files = [
        ("pdf_file", ("doc.pdf", pdf_3page, "application/pdf")),
        ("sig_file", ("sig.png", sig_png, "image/png")),
    ]
    data = {
        "page_number": "1",
        "pos_x_pct": "65",
        "pos_y_pct": "80",
        "sig_width_pt": "120",
    }
    res = client.post("/api/tools/pdf-sign", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc[0].get_images()) >= 1
    doc.close()


def test_scan_cleanup(client, jpg_photo):
    files = [("file", ("scan.jpg", jpg_photo, "image/jpeg"))]
    data = {"mode": "bw", "contrast": "1.8", "brightness": "1.2"}
    res = client.post("/api/tools/scan-cleanup", files=files, data=data)
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 1
    doc.close()


def test_redact(client, pdf_3page):
    initial_doc = fitz.open(stream=pdf_3page, filetype="pdf")
    assert "SECRET_AADHAAR_9999" in initial_doc[1].get_text()
    initial_doc.close()

    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {
        "page_number": "2",
        "x_pct": "5",
        "y_pct": "10",
        "w_pct": "90",
        "h_pct": "10",
        "style": "blackout",
    }
    res = client.post("/api/tools/redact", files=files, data=data)
    assert res.status_code == 200

    redacted_doc = fitz.open(stream=res.content, filetype="pdf")
    assert "SECRET_AADHAAR_9999" not in redacted_doc[1].get_text()
    redacted_doc.close()


# ── STEP 5: Previews ────────────────────────────────────────────────────────

def test_preview_pdf_thumbnails(client, pdf_120page):
    files = [("file", ("large.pdf", pdf_120page, "application/pdf"))]
    data = {"page_start": "1", "page_end": "120", "batch_size": "24"}
    res = client.post("/api/preview/pdf-thumbnails", files=files, data=data)
    assert res.status_code == 200
    res_json = res.json()
    assert res_json["total_pages"] == 120
    assert len(res_json["thumbnails"]) == 24
    assert res_json["thumbnails"][0]["data"].startswith("data:image/jpeg;base64,")


def test_preview_watermark(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"text": "PREVIEW", "page": "1", "color": "gray", "opacity": "0.3"}
    res = client.post("/api/preview/watermark", files=files, data=data)
    assert res.status_code == 200
    b64_str = res.json()["data"].split(",", 1)[1]
    img = Image.open(io.BytesIO(base64.b64decode(b64_str)))
    assert img.size[0] > 0


def test_preview_page_numbers(client, pdf_3page):
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    data = {"position": "bottom-center", "format": "page", "start_num": "1", "margin": "36", "page": "1"}
    res = client.post("/api/preview/page-numbers", files=files, data=data)
    assert res.status_code == 200
    b64_str = res.json()["data"].split(",", 1)[1]
    img = Image.open(io.BytesIO(base64.b64decode(b64_str)))
    assert img.size[0] > 0


def test_preview_doc_page(client, sample_docx):
    files = [("file", ("test.docx", sample_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))]
    data = {"page": "1"}
    res = client.post("/api/preview/doc-page", files=files, data=data)
    assert res.status_code == 200
    res_json = res.json()
    if not res_json.get("available"):
        assert "reason" in res_json
    else:
        assert "data" in res_json
        b64_str = res_json["data"].split(",", 1)[1]
        img = Image.open(io.BytesIO(base64.b64decode(b64_str)))
        assert img.size[0] > 0


# ── STEP 6: My Files ────────────────────────────────────────────────────────

def test_my_files_default_outside_onedrive():
    default_dir = get_default_my_files_dir()
    str_dir = str(default_dir).lower()
    assert "onedrive" not in str_dir, f"Default directory {default_dir} should not be inside OneDrive"


def test_my_files_crud_and_folders(client, jpg_photo):
    # 1. Save small file (< 100 KB) -> lands in "Under 100 KB"
    files = [("file", ("test_photo.jpg", jpg_photo, "image/jpeg"))]
    res_save = client.post("/api/files/save", files=files, data={"folder": "", "filename": "test_photo.jpg"})
    assert res_save.status_code == 200
    save_info = res_save.json()
    assert save_info["folder"] == "Under 100 KB"
    fname = save_info["filename"]

    # 2. List files
    res_list = client.get("/api/files/list", params={"folder": "Under 100 KB"})
    assert res_list.status_code == 200
    files_list = res_list.json()
    assert any(f["name"] == fname for f in files_list)

    # 3. Rename file with unique name to prevent collision across test runs
    import uuid
    new_name = f"renamed_test_photo_{uuid.uuid4().hex[:6]}.jpg"
    res_ren = client.post("/api/files/rename", data={"folder": "Under 100 KB", "old_name": fname, "new_name": new_name})
    assert res_ren.status_code == 200

    # 4. Move file
    res_mov = client.post("/api/files/move", data={"source_folder": "Under 100 KB", "filename": new_name, "target_folder": "Photos"})
    assert res_mov.status_code == 200

    # 5. Download file
    res_down = client.get(f"/api/files/download/Photos/{new_name}")
    assert res_down.status_code == 200
    assert len(res_down.content) == len(jpg_photo)

    # 6. Delete file
    res_del = client.post("/api/files/delete", data={"folder": "Photos", "filename": new_name})
    assert res_del.status_code == 200

    # 7. Open folder endpoint
    res_open = client.post("/api/files/open-folder", data={"folder": "Photos"})
    assert res_open.status_code == 200


# ── STEP 7: User Settings & History ─────────────────────────────────────────

def test_user_settings_and_history(client):
    # Reset history_paused to False before test
    client.post("/api/user/settings", json={"history_paused": False})
    # 1. Get initial settings
    res = client.get("/api/user/settings")
    assert res.status_code == 200
    settings = res.json()
    # Recent files must be OFF by default
    assert settings.get("recent_files_enabled") is False

    # 2. Record tool usage
    res_rec = client.post("/api/user/record-tool", data={"tool_id": "pdf-merge"})
    assert res_rec.status_code == 200
    assert any(t["id"] == "pdf-merge" for t in res_rec.json().get("recent_tools", []))

    # 3. Toggle favorite
    # Ensure tool starts not in favorites
    init_favs = client.get("/api/user/settings").json().get("favorites", [])
    if "pdf-merge" in init_favs:
        client.post("/api/user/toggle-favorite", data={"tool_id": "pdf-merge"})
    res_fav = client.post("/api/user/toggle-favorite", data={"tool_id": "pdf-merge"})
    assert res_fav.status_code == 200
    assert "pdf-merge" in res_fav.json().get("favorites", [])

    # 4. Pause history
    res_pause = client.post("/api/user/settings", json={"history_paused": True})
    assert res_pause.status_code == 200
    assert res_pause.json().get("settings", {}).get("history_paused") is True

    # Recording while paused should not add new tools
    client.post("/api/user/record-tool", data={"tool_id": "pdf-watermark"})
    res_chk = client.get("/api/user/settings")
    recent_ids = [t["id"] for t in res_chk.json().get("recent_tools", [])]
    assert "pdf-watermark" not in recent_ids

    # 5. Clear history
    res_clr = client.post("/api/user/clear-history")
    assert res_clr.status_code == 200
    assert len(res_clr.json().get("recent_tools", [])) == 0

    # Restore history paused = False
    client.post("/api/user/settings", json={"history_paused": False})


# ── STEP 8: Presets ─────────────────────────────────────────────────────────

def test_presets_list_and_custom(client):
    # 1. List presets (all 12 built-in presets present)
    res = client.get("/api/presets")
    assert res.status_code == 200
    presets = res.json()
    assert len(presets) >= 12

    # 2. Add custom preset
    custom_preset = {
        "id": "custom-smoke-test-preset",
        "portal": "Custom State Portal",
        "name": "Candidate Application Photo",
        "category": "photo",
        "max_kb": 60,
        "min_kb": 20,
        "width_px": 350,
        "height_px": 450,
        "dpi": 150,
        "format": "jpg",
        "note": "Custom smoke test preset",
    }
    updated_list = presets + [custom_preset]
    res_save = client.post("/api/presets", json=updated_list)
    assert res_save.status_code == 200

    # 3. Retrieve back
    res_back = client.get("/api/presets")
    assert res_back.status_code == 200
    ids = [p["id"] for p in res_back.json()]
    assert "custom-smoke-test-preset" in ids


# ── STEP 9: Security Checks ─────────────────────────────────────────────────

def test_security_tokens_and_hosts(client):
    # 1. Session token enforcement
    os.environ["LTK_ENFORCE_TOKEN_TEST"] = "1"
    try:
        # Without token -> 403 Forbidden
        res_no_token = client.get("/api/user/settings")
        assert res_no_token.status_code == 403
        assert "Invalid or missing session token" in res_no_token.json().get("detail", "")

        # With header token -> 200 OK
        res_with_hdr = client.get("/api/user/settings", headers={"x-session-token": LAUNCH_TOKEN})
        assert res_with_hdr.status_code == 200

        # With query token -> 200 OK
        res_with_qry = client.get(f"/api/user/settings?token={LAUNCH_TOKEN}")
        assert res_with_qry.status_code == 200
    finally:
        os.environ.pop("LTK_ENFORCE_TOKEN_TEST", None)

    # 2. Host header security
    # Allowed hosts
    res_local = client.get("/api/health", headers={"Host": "localhost:8791"})
    assert res_local.status_code == 200
    res_ip = client.get("/api/health", headers={"Host": "127.0.0.1:8791"})
    assert res_ip.status_code == 200

    # Rejected host
    res_evil = client.get("/api/health", headers={"Host": "attacker-domain.com"})
    assert res_evil.status_code == 403


def test_error_handling_readable_404_405_422(client):
    # 1. 404 Not Found JSON on unknown /api/*
    res_404 = client.get("/api/unknown-tool-route-xyz")
    assert res_404.status_code == 404
    assert res_404.json() == {"detail": "Not found"}

    # 2. 405 Method Not Allowed JSON
    res_405 = client.get("/api/pdf/merge")
    assert res_405.status_code == 405
    assert res_405.json() == {"detail": "Method not allowed"}

    # 3. Readable 422 validation error
    res_422 = client.post("/api/files/save")  # Missing required 'file' field
    assert res_422.status_code == 422
    err_json = res_422.json()
    assert "detail" in err_json
    assert "file" in err_json["detail"].lower() or "missing" in err_json["detail"].lower()


# ── STEP 10: Static Assets Check (Zero External CDNs) ───────────────────────

def test_static_frontend_no_external_cdns():
    """Scan all frontend HTML/CSS/JS files and assert NO external resource URLs."""
    frontend_dir = Path("frontend")
    assert frontend_dir.exists()

    # Pattern for external resource loads: src/href/url() pointing to external schemas
    resource_pattern = re.compile(
        r"""(?:src|href)=['"](https?://[^'"]+)['"]|url\(['"]?(https?://[^)'"]+)['"]?\)|@import\s+['"]?(https?://[^)'"]+)['"]?""",
        re.IGNORECASE,
    )

    forbidden_domains = [
        "googleapis.com",
        "gstatic.com",
        "cdnjs.cloudflare.com",
        "unpkg.com",
        "jsdelivr.net",
        "cdn.",
    ]

    violations = []
    for fpath in frontend_dir.rglob("*"):
        if fpath.suffix.lower() in (".html", ".css", ".js"):
            content = fpath.read_text(encoding="utf-8", errors="ignore")
            for line_no, line in enumerate(content.splitlines(), 1):
                # Skip informative comments or tool download links (e.g. libreoffice install link in config)
                if "depInstallUrl" in line or "https://www.libreoffice.org" in line:
                    continue
                for m in resource_pattern.finditer(line):
                    url = m.group(1) or m.group(2) or m.group(3)
                    if any(domain in url.lower() for domain in forbidden_domains):
                        violations.append((str(fpath), line_no, url))

    assert not violations, f"Found external CDN or font references in offline frontend: {violations}"


# ── STEP 11: Regression Tests for Reported Browser Issues ────────────────────

def test_regression_remove_pages_typed_and_bounds(client, pdf_3page):
    """Verify remove-pages handles typed ranges, reverse ranges, empty checks, and bounds validation."""
    # 1. Normal removal
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/remove-pages", files=files, data={"pages": "2"})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 2
    doc.close()

    # 2. Reverse range (3-2)
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/remove-pages", files=files, data={"pages": "3-2"})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 1
    doc.close()

    # 3. Out-of-bounds error with exact page count detail
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/remove-pages", files=files, data={"pages": "15"})
    assert res.status_code == 400
    assert "Page 15 does not exist; this PDF has 3 pages." in res.json().get("detail", "")

    # 4. Cannot remove all pages error
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/remove-pages", files=files, data={"pages": "1-3"})
    assert res.status_code == 400
    assert "Cannot remove all pages" in res.json().get("detail", "")

    # 5. Empty pages error
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/remove-pages", files=files, data={"pages": "   "})
    assert res.status_code == 400
    assert "specify pages to remove" in res.json().get("detail", "").lower()


def test_regression_rotate_pdf_angle_range_and_offsets(client, pdf_3page):
    """Verify rotate PDF supports global angle, page subset, blank/all, and per-page visualizer offsets."""
    # 1. Global angle + page subset ("1, 3")
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/rotate", files=files, data={"angle": "90", "pages": "1, 3"})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert doc[0].rotation == 90
    assert doc[1].rotation == 0
    assert doc[2].rotation == 90
    doc.close()

    # 2. Blank / "all" pages
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/rotate", files=files, data={"angle": "180", "pages": ""})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert doc[0].rotation == 180
    assert doc[1].rotation == 180
    assert doc[2].rotation == 180
    doc.close()

    # 3. Per-page offsets from visualizer (page 2 +90 deg offset)
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    offsets = json.dumps({"2": 90})
    res = client.post("/api/pdf/rotate", files=files, data={"angle": "0", "pages": "all", "page_rotations": offsets})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert doc[0].rotation == 0
    assert doc[1].rotation == 90
    assert doc[2].rotation == 0
    doc.close()


def test_regression_extract_pdf_validation(client, pdf_3page):
    """Verify extract-pages validates inputs and handles reversed ranges cleanly without crashing."""
    # 1. Valid extraction with spaces
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/extract", files=files, data={"pages": "  1 , 3  "})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 2
    doc.close()

    # 2. Reversed range (3-1)
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/extract", files=files, data={"pages": "3-1"})
    assert res.status_code == 200
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert len(doc) == 3
    doc.close()

    # 3. Out-of-range error
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/extract", files=files, data={"pages": "99"})
    assert res.status_code == 400
    assert "Page 99 does not exist; this PDF has 3 pages." in res.json().get("detail", "")

    # 4. Invalid input string
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/extract", files=files, data={"pages": "abc"})
    assert res.status_code == 400
    assert "invalid" in res.json().get("detail", "").lower() or "page" in res.json().get("detail", "").lower()


def test_regression_pdf_thumbnails_request_batching(client, pdf_120page):
    """Verify batching allows fetching 24 thumbnails at a time with correct total_pages."""
    files = [("file", ("doc120.pdf", pdf_120page, "application/pdf"))]
    data = {"max_pages": "24", "start_page": "1", "dpi": "36"}
    res = client.post("/api/preview/pdf-thumbnails", files=files, data=data)
    assert res.status_code == 200
    body = res.json()
    assert body["total_pages"] == 120
    assert len(body["thumbnails"]) == 24
    # Check start page is 1
    assert body["thumbnails"][0]["page"] == 1
    assert body["thumbnails"][-1]["page"] == 24


def test_regression_pdf_compress_ladder_and_breakdown(client, pdf_3page):
    """Verify PDF compression uses transparent 3-step ladder and returns X-Compression-Report header."""
    files = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res = client.post("/api/pdf/compress", files=files, data={"target_kb": "100", "quality": "medium"})
    assert res.status_code == 200
    assert "X-Compression-Report" in res.headers

    report_str = urllib.parse.unquote(res.headers["X-Compression-Report"])
    report = json.loads(report_str)
    assert "original_size" in report
    assert "final_size" in report
    assert "target_size" in report
    assert "breakdown" in report
    assert "steps_saved" in report
    assert report["steps_saved"]["lossless_cleanup"] >= 0
    assert report["steps_saved"]["image_compression"] >= 0
    assert report["breakdown"]["images"] > 0

    # Also test estimate endpoint returns synchronized report
    files_est = [("file", ("doc.pdf", pdf_3page, "application/pdf"))]
    res_est = client.post("/api/pdf/estimate-size", files=files_est, data={"target_kb": "100", "quality": "medium"})
    assert res_est.status_code == 200
    est_json = res_est.json()
    assert "estimated_bytes" in est_json
    assert "report" in est_json
    assert est_json["report"]["steps_saved"]["lossless_cleanup"] >= 0

